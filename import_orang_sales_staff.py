"""
Import Orange Protection sales staff from SAP HANA (4B-ORANG_LIVE),
create Django users with sales-staff role, generate passwords,
assign region/zone/territory, and export a complete Excel sheet.

Filtered designations (OHPS position IDs):
  9  → Dept Product Leader      → DEPUTY PRODUCT LEADER
  10 → Dpt.Reg Sales Leader     → DEPUTY REGIONAL SALES LEADER
  19 → Product Leader           → PRODUCT LEADER
  31 → Sr. Product Leader       → SENIOR PRODUCT LEADER
  39 → Regional SL              → REGIONAL SALES LEADER
  45 → Dept.Reg.Sale.Leader     → DEPUTY REGIONAL SALES LEADER
"""

import base64, hashlib, os, random, string, sys
import pymysql
from hdbcli import dbapi
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ── Config ───────────────────────────────────────────────────────────────────

HANA = dict(
    address='fourb.vdc.services', port=30015,
    user='SYSTEM', password='S@pHFP21*',
    encrypt=True, sslValidateCertificate=False,
)
HANA_SCHEMA = '4B-ORANG_LIVE'

DB = dict(host='127.0.0.1', port=3306, user='agritech',
          password='pnyABiGnzWjXasnw', database='agritech', charset='utf8mb4')

OUT_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                        'orang_sales_staff_passwords.xlsx')

# SAP OHPS posID → (django_designation_id, label)
POSITION_MAP = {
    9:  (255, 'Deputy Product Leader'),
    10: (252, 'Deputy Regional Sales Leader'),
    19: (254, 'Product Leader'),
    31: (253, 'Senior Product Leader'),
    39: (134, 'Regional Sales Leader'),
    45: (252, 'Deputy Regional Sales Leader'),
}
TARGET_POS_IDS = tuple(POSITION_MAP.keys())

ORANG_COMPANY_ID = 2   # fieldadvisoryservice_company id for 4B-ORANG

# ── Password ─────────────────────────────────────────────────────────────────

def _make_password(first_name):
    """FirstName + 3 digits + 2 uppercase  →  Ahmad847KP"""
    name = (first_name or 'Emp').split()[0].capitalize()
    digits = ''.join(random.choices(string.digits, k=3))
    upper  = ''.join(random.choices(string.ascii_uppercase, k=2))
    return f"{name}{digits}{upper}"


def _django_hash(plain, iterations=870000):
    salt = base64.b64encode(os.urandom(12)).decode('ascii').rstrip('=')
    dk   = hashlib.pbkdf2_hmac('sha256', plain.encode(), salt.encode(), iterations)
    return f"pbkdf2_sha256${iterations}${salt}${base64.b64encode(dk).decode()}"


# ── HANA fetch ────────────────────────────────────────────────────────────────

def fetch_hana_employees():
    conn = dbapi.connect(**HANA)
    cur  = conn.cursor()
    cur.execute(f'SET SCHEMA "{HANA_SCHEMA}"')

    placeholders = ','.join(str(p) for p in TARGET_POS_IDS)
    cur.execute(f"""
        SELECT
            e."empID",
            e."firstName",
            e."middleName",
            e."lastName",
            e."email",
            e."mobile",
            e."homeTel",
            e."position",
            e."U_TERR"
        FROM "OHEM" e
        WHERE e."Active" = 'Y'
          AND e."position" IN ({placeholders})
        ORDER BY e."empID"
    """)
    cols = [d[0] for d in cur.description]
    rows = [dict(zip(cols, r)) for r in cur.fetchall()]
    cur.close()
    conn.close()
    return rows


# ── DB helpers ────────────────────────────────────────────────────────────────

def get_hana_territory_map(conn):
    """hana_territory_id → (territory_id, territory_name, zone_name, region_name)"""
    cur = conn.cursor()
    cur.execute("""
        SELECT t.hana_territory_id, t.id, t.name, z.name, r.name
        FROM fieldadvisoryservice_territory t
        JOIN fieldadvisoryservice_zone z ON z.id = t.zone_id
        JOIN fieldadvisoryservice_region r ON r.id = z.region_id
        JOIN fieldadvisoryservice_company c ON c.id = t.company_id
        WHERE c.Company_name = '4B-ORANG'
          AND t.hana_territory_id IS NOT NULL
    """)
    mapping = {}
    for hana_id, t_id, t_name, z_name, r_name in cur.fetchall():
        mapping[str(hana_id)] = (t_id, t_name, z_name, r_name)
    cur.close()
    return mapping


def get_or_create_user(conn, emp):
    """Return (user_id, created). Creates user if not exists by email."""
    cur = conn.cursor()
    email    = (emp['email'] or '').strip().lower() or f"emp{emp['empID']}@4bgroup.com"
    username = email.split('@')[0]
    fname    = (emp['firstName']  or '').strip() or 'Employee'
    lname    = (emp['lastName']   or '').strip() or str(emp['empID'])

    cur.execute("SELECT id FROM accounts_user WHERE email = %s", (email,))
    row = cur.fetchone()
    if row:
        user_id = row[0]
        # Update flags
        cur.execute("""
            UPDATE accounts_user
            SET is_sales_staff=1, is_dealer=0, role_id=3, is_active=1, company_id=%s
            WHERE id=%s
        """, (ORANG_COMPANY_ID, user_id))
        conn.commit()
        cur.close()
        return user_id, False

    # Check username uniqueness
    base_uname = username
    suffix = 1
    while True:
        cur.execute("SELECT id FROM accounts_user WHERE username=%s", (username,))
        if not cur.fetchone():
            break
        username = f"{base_uname}{suffix}"
        suffix += 1

    cur.execute("""
        INSERT INTO accounts_user
            (username, email, first_name, last_name, is_active,
             is_staff, is_superuser, is_sales_staff, is_dealer,
             role_id, company_id, date_joined, password)
        VALUES (%s,%s,%s,%s,1, 0,0, 1,0, 3,%s, NOW(), '')
    """, (username, email, fname, lname, ORANG_COMPANY_ID))
    conn.commit()
    user_id = cur.lastrowid
    cur.close()
    return user_id, True


def get_or_create_profile(conn, user_id, emp, desig_id):
    """Return (profile_id, created)."""
    employee_code = str(emp['empID'])
    mobile = (emp['mobile'] or emp['homeTel'] or '').strip() or 'N/A'

    cur = conn.cursor()
    cur.execute("SELECT id FROM accounts_salesstaffprofile WHERE user_id=%s", (user_id,))
    row = cur.fetchone()
    if row:
        profile_id = row[0]
        cur.execute("""
            UPDATE accounts_salesstaffprofile
            SET employee_code=%s, designation_id=%s, is_vacant=0
            WHERE id=%s
        """, (employee_code, desig_id, profile_id))
        conn.commit()
        cur.close()
        return profile_id, False

    cur.execute("""
        INSERT INTO accounts_salesstaffprofile
            (user_id, employee_code, designation_id, phone_number,
             address, is_vacant, manager_id,
             sick_leave_quota, casual_leave_quota, others_leave_quota)
        VALUES (%s,%s,%s,%s, '', 0, NULL, 0, 0, 0)
    """, (user_id, employee_code, desig_id, mobile))
    conn.commit()
    profile_id = cur.lastrowid
    cur.close()
    return profile_id, True


def assign_territory(conn, profile_id, territory_id):
    """Assign territory (and its zone/region) to profile via M2M."""
    cur = conn.cursor()

    # Get zone and region for this territory
    cur.execute("""
        SELECT t.id, t.zone_id, z.region_id
        FROM fieldadvisoryservice_territory t
        JOIN fieldadvisoryservice_zone z ON z.id = t.zone_id
        WHERE t.id = %s
    """, (territory_id,))
    row = cur.fetchone()
    if not row:
        cur.close()
        return
    _, zone_id, region_id = row

    for table, col, val in [
        ('accounts_salesstaffprofile_territories', 'territory_id', territory_id),
        ('accounts_salesstaffprofile_zones',       'zone_id',       zone_id),
        ('accounts_salesstaffprofile_regions',     'region_id',     region_id),
    ]:
        cur.execute(
            f"SELECT 1 FROM {table} WHERE salesstaffprofile_id=%s AND {col}=%s",
            (profile_id, val)
        )
        if not cur.fetchone():
            cur.execute(
                f"INSERT INTO {table} (salesstaffprofile_id, {col}) VALUES (%s,%s)",
                (profile_id, val)
            )
    conn.commit()
    cur.close()


def link_company(conn, profile_id, employee_code):
    """Ensure SalesStaffCompany row exists for 4B-ORANG."""
    cur = conn.cursor()
    cur.execute("""
        SELECT 1 FROM accounts_salesstaffcompany
        WHERE sales_profile_id=%s AND company_id=%s
    """, (profile_id, ORANG_COMPANY_ID))
    if not cur.fetchone():
        cur.execute("""
            INSERT INTO accounts_salesstaffcompany
                (sales_profile_id, company_id, employee_code, is_primary, is_active)
            VALUES (%s,%s,%s,1,1)
        """, (profile_id, ORANG_COMPANY_ID, employee_code))
        conn.commit()
    cur.close()


def save_password(conn, user_id, hashed):
    cur = conn.cursor()
    cur.execute("UPDATE accounts_user SET password=%s WHERE id=%s", (hashed, user_id))
    conn.commit()
    cur.close()


def get_username(conn, user_id):
    cur = conn.cursor()
    cur.execute("SELECT username FROM accounts_user WHERE id=%s", (user_id,))
    row = cur.fetchone()
    cur.close()
    return row[0] if row else ''


# ── Excel ─────────────────────────────────────────────────────────────────────

def build_excel(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Orange Sales Staff'

    hdr_font = Font(bold=True, color='FFFFFF', size=11)
    hdr_fill = PatternFill('solid', fgColor='C45911')
    alt_fill = PatternFill('solid', fgColor='FCE4D6')
    center   = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left     = Alignment(horizontal='left',   vertical='center', wrap_text=True)
    thin     = Side(style='thin', color='CCCCCC')
    border   = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers    = ['#', 'Emp Code', 'Full Name', 'Username', 'Email',
                  'Phone', 'Password', 'Designation',
                  'Region', 'Zone', 'Territory']
    col_widths = [5,    11,          26,           22,         32,
                  18,    18,           28,
                  24,      22,      26]

    for col, (h, w) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = hdr_font; cell.fill = hdr_fill
        cell.alignment = center; cell.border = border
        ws.column_dimensions[cell.column_letter].width = w

    ws.row_dimensions[1].height = 22
    ws.freeze_panes = 'A2'

    for i, row in enumerate(rows, start=2):
        fill = alt_fill if i % 2 == 0 else None
        vals = [i-1, row['emp_code'], row['name'], row['username'], row['email'],
                row['phone'], row['password'], row['designation'],
                row['region'], row['zone'], row['territory']]
        for col, val in enumerate(vals, start=1):
            cell = ws.cell(row=i, column=col, value=val)
            cell.border    = border
            cell.alignment = center if col <= 2 else left
            if fill: cell.fill = fill
        ws.row_dimensions[i].height = 18

    ws.cell(row=len(rows)+2, column=1,
            value=f'Total: {len(rows)} employees').font = Font(bold=True, italic=True)
    wb.save(OUT_PATH)


# ── Main ─────────────────────────────────────────────────────────────────────

def main():
    print('Connecting to SAP HANA ...')
    employees = fetch_hana_employees()
    print(f'Found {len(employees)} matching employees in HANA ({HANA_SCHEMA})')

    if not employees:
        print('No employees found for the selected designations. Exiting.')
        sys.exit(0)

    print('Connecting to Django DB ...')
    conn = pymysql.connect(**DB)
    terr_map = get_hana_territory_map(conn)
    print(f'Territory map loaded: {len(terr_map)} HANA territories')

    created_count = updated_count = 0
    rows = []

    for emp in employees:
        emp_id    = emp['empID']
        pos_id    = emp['position']
        desig_id, desig_label = POSITION_MAP.get(pos_id, (140, 'Territory Sales Officer'))
        fname     = (emp['firstName']  or '').strip() or 'Employee'
        lname     = (emp['lastName']   or '').strip() or str(emp_id)
        mobile    = (emp['mobile'] or emp['homeTel'] or '').strip()
        email     = (emp['email'] or '').strip().lower() or f"emp{emp_id}@4bgroup.com"
        u_terr    = str(emp['U_TERR']).strip() if emp['U_TERR'] else None

        # Territory lookup
        terr_info   = terr_map.get(u_terr) if u_terr else None
        territory_id   = terr_info[0] if terr_info else None
        territory_name = terr_info[1] if terr_info else '—'
        zone_name      = terr_info[2] if terr_info else '—'
        region_name    = terr_info[3] if terr_info else '—'

        # Create / update user
        user_id, created = get_or_create_user(conn, emp)
        if created:
            created_count += 1
        else:
            updated_count += 1

        # Create / update SalesStaffProfile
        profile_id, _ = get_or_create_profile(conn, user_id, emp, desig_id)

        # Link to Orange company
        link_company(conn, profile_id, str(emp_id))

        # Assign territory M2M
        if territory_id:
            assign_territory(conn, profile_id, territory_id)

        # Generate + save password
        plain  = _make_password(fname)
        hashed = _django_hash(plain)
        save_password(conn, user_id, hashed)

        username = get_username(conn, user_id)

        full_name = f"{fname} {lname}".strip()
        print(f"  {'NEW' if created else 'UPD'}  {emp_id:<8} {full_name:<28} {desig_label}")

        rows.append({
            'emp_code':    str(emp_id),
            'name':        full_name,
            'username':    username,
            'email':       email,
            'phone':       mobile or '—',
            'password':    plain,
            'designation': desig_label,
            'region':      region_name,
            'zone':        zone_name,
            'territory':   territory_name,
        })

    conn.close()

    rows.sort(key=lambda r: r['name'].lower())
    build_excel(rows)

    print()
    print('='*55)
    print(f'  Created : {created_count}')
    print(f'  Updated : {updated_count}')
    print(f'  Total   : {len(rows)}')
    print(f'  Sheet   : {OUT_PATH}')
    print('='*55)


if __name__ == '__main__':
    main()
