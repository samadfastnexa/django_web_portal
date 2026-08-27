"""
Standalone script — no Django required.
Generates passwords for Orange (4B-ORANG) sales staff, saves hashed
passwords to DB, and exports an Excel sheet with region/zone/territory.
"""
import base64, csv, hashlib, os, random, string, sys
import pymysql
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

DB = dict(host='127.0.0.1', port=3306, user='agritech',
          password='pnyABiGnzWjXasnw', database='agritech', charset='utf8mb4')

OUT_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                        'orang_employees_passwords.xlsx')


# ── Password helpers ────────────────────────────────────────────────────────

def _random_suffix(n=6):
    return ''.join(random.choices(string.ascii_uppercase + string.digits, k=n))


def _plain_password(first_name):
    return f"{(first_name or '').replace(' ', '')}-{_random_suffix()}"


def _django_hash(plain, iterations=870000):
    """Produce a Django pbkdf2_sha256 hash compatible with check_password()."""
    salt = base64.b64encode(os.urandom(12)).decode('ascii').rstrip('=')
    dk = hashlib.pbkdf2_hmac('sha256', plain.encode(), salt.encode(), iterations)
    return f"pbkdf2_sha256${iterations}${salt}${base64.b64encode(dk).decode()}"


# ── DB queries ───────────────────────────────────────────────────────────────

def fetch_orang_employees(conn):
    cur = conn.cursor()
    cur.execute("""
        SELECT
            sp.id,
            sp.employee_code,
            u.id         AS user_id,
            u.first_name,
            u.last_name,
            u.username,
            u.email,
            u.phone_number,
            d.name       AS designation
        FROM accounts_salesstaffprofile sp
        JOIN accounts_user u ON u.id = sp.user_id
        LEFT JOIN accounts_designationmodel d ON d.id = sp.designation_id
        WHERE sp.employee_code LIKE '4B-ORANG%%'
          AND sp.is_vacant = 0
          AND u.id IS NOT NULL
        ORDER BY u.first_name, u.last_name
    """)
    cols = [c[0] for c in cur.description]
    rows = [dict(zip(cols, r)) for r in cur.fetchall()]
    cur.close()
    return rows


def fetch_territories(conn, profile_ids):
    """Return {profile_id: {regions, zones, territories}}"""
    data = {pid: {'regions': [], 'zones': [], 'territories': []} for pid in profile_ids}
    if not profile_ids:
        return data
    ids_sql = ','.join(str(i) for i in profile_ids)
    cur = conn.cursor()

    cur.execute(f"""
        SELECT m.salesstaffprofile_id, r.name
        FROM accounts_salesstaffprofile_regions m
        JOIN FieldAdvisoryService_region r ON r.id = m.region_id
        WHERE m.salesstaffprofile_id IN ({ids_sql})
    """)
    for pid, name in cur.fetchall():
        if pid in data:
            data[pid]['regions'].append(name)

    cur.execute(f"""
        SELECT m.salesstaffprofile_id, z.name
        FROM accounts_salesstaffprofile_zones m
        JOIN FieldAdvisoryService_zone z ON z.id = m.zone_id
        WHERE m.salesstaffprofile_id IN ({ids_sql})
    """)
    for pid, name in cur.fetchall():
        if pid in data:
            data[pid]['zones'].append(name)

    cur.execute(f"""
        SELECT m.salesstaffprofile_id, t.name
        FROM accounts_salesstaffprofile_territories m
        JOIN FieldAdvisoryService_territory t ON t.id = m.territory_id
        WHERE m.salesstaffprofile_id IN ({ids_sql})
    """)
    for pid, name in cur.fetchall():
        if pid in data:
            data[pid]['territories'].append(name)

    cur.close()
    return data


def save_passwords(conn, updates):
    """updates: list of (user_id, hashed_password)"""
    cur = conn.cursor()
    cur.executemany(
        "UPDATE accounts_user SET password = %s WHERE id = %s",
        [(h, uid) for uid, h in updates]
    )
    conn.commit()
    cur.close()


# ── Excel builder ────────────────────────────────────────────────────────────

def build_excel(rows, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Orange Employees'

    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill('solid', fgColor='C45911')   # orange
    alt_fill    = PatternFill('solid', fgColor='FCE4D6')   # light orange
    center      = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left        = Alignment(horizontal='left',   vertical='center', wrap_text=True)
    thin        = Side(style='thin', color='CCCCCC')
    border      = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers    = ['#', 'Emp Code', 'Full Name', 'Username', 'Email',
                  'Phone', 'Password', 'Designation', 'Regions', 'Zones', 'Territories']
    col_widths = [5,    14,          24,           22,         32,
                  16,    20,          22,             26,        26,      34]

    for col, (h, w) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font, cell.fill, cell.alignment, cell.border = header_font, header_fill, center, border
        ws.column_dimensions[cell.column_letter].width = w

    ws.row_dimensions[1].height = 22
    ws.freeze_panes = 'A2'

    for i, row in enumerate(rows, start=2):
        fill = alt_fill if i % 2 == 0 else None
        vals = [i - 1, row['emp_code'], row['name'], row['username'], row['email'],
                row['phone'], row['password'], row['designation'],
                row['regions'], row['zones'], row['territories']]
        for col, val in enumerate(vals, start=1):
            cell = ws.cell(row=i, column=col, value=val)
            cell.border    = border
            cell.alignment = center if col <= 2 else left
            if fill:
                cell.fill = fill
        ws.row_dimensions[i].height = 18

    ws.cell(row=len(rows) + 2, column=1,
            value=f'Total: {len(rows)} employees').font = Font(bold=True, italic=True)

    wb.save(out_path)


# ── Main ─────────────────────────────────────────────────────────────────────

def main():
    conn = pymysql.connect(**DB)

    employees = fetch_orang_employees(conn)
    print(f'Orange employees found: {len(employees)}')

    if not employees:
        print('No Orange employees in DB.')
        conn.close()
        sys.exit(0)

    profile_ids = [e['id'] for e in employees]
    terr_map    = fetch_territories(conn, profile_ids)

    updates = []
    rows    = []

    for emp in employees:
        plain    = _plain_password(emp['first_name'])
        hashed   = _django_hash(plain)
        updates.append((emp['user_id'], hashed))

        t = terr_map.get(emp['id'], {})
        rows.append({
            'emp_code':    emp['employee_code'] or '—',
            'name':        f"{emp['first_name']} {emp['last_name']}".strip(),
            'username':    emp['username'] or '—',
            'email':       emp['email'] or '—',
            'phone':       emp['phone_number'] or '—',
            'password':    plain,
            'designation': emp['designation'] or '—',
            'regions':     ', '.join(t.get('regions', []))     or '—',
            'zones':       ', '.join(t.get('zones', []))       or '—',
            'territories': ', '.join(t.get('territories', [])) or '—',
        })

    print(f'Saving {len(updates)} password updates to DB ...')
    save_passwords(conn, updates)
    conn.close()
    print('Done.')

    build_excel(rows, OUT_PATH)
    print(f'Saved: {OUT_PATH}  ({len(rows)} rows)')


if __name__ == '__main__':
    main()
