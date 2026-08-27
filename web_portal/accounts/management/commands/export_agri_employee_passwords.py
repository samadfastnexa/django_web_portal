import csv
import os
from django.core.management.base import BaseCommand
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


class Command(BaseCommand):
    help = "Export Agri employees with passwords and their region/zone/territory assignments to Excel"

    def add_arguments(self, parser):
        parser.add_argument(
            '--out',
            type=str,
            default='agri_employees_passwords.xlsx',
            help='Output Excel file path (default: agri_employees_passwords.xlsx)',
        )
        parser.add_argument(
            '--csv',
            type=str,
            default=None,
            help='Path to Agri_Employee.csv (auto-detected if not given)',
        )

    def handle(self, *args, **options):
        from accounts.models import SalesStaffProfile

        csv_path = options['csv'] or self._find_csv()
        if not csv_path or not os.path.exists(csv_path):
            self.stderr.write(self.style.ERROR(f'CSV not found: {csv_path}'))
            return

        out_path = options['out']

        # Read CSV into dict keyed by emp_code
        self.stdout.write(f'Reading CSV: {csv_path}')
        csv_data = {}
        with open(csv_path, encoding='utf-8-sig') as f:
            for row in csv.DictReader(f):
                code = (row.get('EmpCode') or '').strip()
                if code:
                    csv_data[code] = row

        self.stdout.write(f'  {len(csv_data)} employees in CSV')

        # Build lookup: employee_code -> profile
        profiles = (
            SalesStaffProfile.objects
            .filter(employee_code__isnull=False)
            .prefetch_related('regions', 'zones', 'territories')
        )
        profile_map = {p.employee_code.strip(): p for p in profiles if p.employee_code}
        self.stdout.write(f'  {len(profile_map)} profiles in DB')

        # Build rows
        rows = []
        for emp_code, emp in csv_data.items():
            profile = profile_map.get(emp_code)

            if profile:
                regions = ', '.join(r.name for r in profile.regions.all()) or '—'
                zones   = ', '.join(z.name for z in profile.zones.all())   or '—'
                terrs   = ', '.join(t.name for t in profile.territories.all()) or '—'
                designation = getattr(profile, 'designation', None)
                desig_name = str(designation) if designation else '—'
            else:
                regions = zones = terrs = desig_name = '—'

            rows.append({
                'EmpCode':     emp_code,
                'Full Name':   emp.get('FullName', '').strip(),
                'Username':    emp.get('Username', '').strip(),
                'Email':       emp.get('Email', '').strip(),
                'Phone':       emp.get('PhoneNumber', '').strip(),
                'Password':    emp.get('Password', '').strip(),
                'Designation': desig_name,
                'Regions':     regions,
                'Zones':       zones,
                'Territories': terrs,
            })

        rows.sort(key=lambda r: r['Full Name'].lower())

        # Build Excel
        wb = Workbook()
        ws = wb.active
        ws.title = 'Agri Employees'

        # Styles
        header_font   = Font(bold=True, color='FFFFFF', size=11)
        header_fill   = PatternFill('solid', fgColor='1F6B3A')  # dark green
        alt_fill      = PatternFill('solid', fgColor='EAF4EC')  # light green
        center        = Alignment(horizontal='center', vertical='center', wrap_text=True)
        left          = Alignment(horizontal='left',   vertical='center', wrap_text=True)
        thin          = Side(style='thin', color='CCCCCC')
        border        = Border(left=thin, right=thin, top=thin, bottom=thin)

        headers = ['EmpCode', 'Full Name', 'Username', 'Email', 'Phone',
                   'Password', 'Designation', 'Regions', 'Zones', 'Territories']
        col_widths = [10, 24, 22, 32, 16, 20, 22, 22, 22, 30]

        # Header row
        for col, (h, w) in enumerate(zip(headers, col_widths), start=1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font      = header_font
            cell.fill      = header_fill
            cell.alignment = center
            cell.border    = border
            ws.column_dimensions[cell.column_letter].width = w

        ws.row_dimensions[1].height = 22
        ws.freeze_panes = 'A2'

        # Data rows
        for i, row in enumerate(rows, start=2):
            fill = alt_fill if i % 2 == 0 else None
            for col, key in enumerate(headers, start=1):
                cell = ws.cell(row=i, column=col, value=row[key])
                cell.border    = border
                cell.alignment = center if col == 1 else left
                if fill:
                    cell.fill = fill
            ws.row_dimensions[i].height = 18

        # Summary in last row
        summary_row = len(rows) + 2
        ws.cell(row=summary_row, column=1, value=f'Total: {len(rows)} employees').font = Font(bold=True, italic=True)

        wb.save(out_path)
        self.stdout.write(self.style.SUCCESS(f'Saved: {out_path}  ({len(rows)} rows)'))

    def _find_csv(self):
        candidates = [
            os.path.join(os.path.dirname(__file__), '..', '..', '..', 'csv_data', 'Agri_Employee.csv'),
            os.path.join(os.path.dirname(__file__), '..', '..', '..', 'data',     'Agri_Employee.csv'),
        ]
        for p in candidates:
            resolved = os.path.normpath(p)
            if os.path.exists(resolved):
                return resolved
        return None
