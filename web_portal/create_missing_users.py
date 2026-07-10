"""
Run with:
    python manage.py shell < create_missing_users.py
"""
import random
import string
import os
from django.db import transaction
from accounts.models import User, SalesStaffProfile, Role
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

SALES_STAFF_ROLE_ID = 3   # confirmed from DB

# ── Password builder (matches generate_passwords.py logic) ────────────────────
def _random_suffix(length=6):
    return ''.join(random.choices(string.ascii_uppercase + string.digits, k=length))

def build_password(first_name):
    name_part = (first_name or '').replace(' ', '')
    return f"{name_part}-{_random_suffix()}"

# ── New users to create (deduplicated — Manzoor Amir 2010 appears twice) ──────
USERS = [
    # (first_name,       last_name,        employee_code, email)
    ('Muhammad',         'Abbas',           '1093', 'muhammad.abbas@4bgroup.com'),
    ('Asghar',           'Ali',             '1140', 'asghar.nukrach@4bgroup.com'),
    ('Noor',             'Hassan',          '1968', 'noor.hassan@4bgroup.com'),
    ('Abdul',            'Waheed',          '1974', 'm.abdulwaheed@4bgroup.com'),
    ('Shahzad',          'Hussain',         '2006', 'shahzad.hussain@4bgroup.com'),
    ('Amir',             'Manzoor',         '2010', 'amir.manzoor@4bgroup.com'),
    ('Muhammad',         'Mairaj',          '2063', 'mairaj.muhammad@4bgroup.com'),
    ('Noman',            'Rasheed',         '2080', 'noman.rasheed@4bgroup.com'),
    ('Muhammad',         'Muaz',            '2090', 'muhammad.muaz@4bgroup.com'),
    ('Ishtiaq',          'Alam',            '2099', 'ishtiaq.alam@4bgroup.com'),
    ('Bilal',            'Wahid',           '2100', 'bilal.wahid@4bgroup.com'),
    ('Mehtab',           'Shah',            '2102', 'mehtab.shah@4bgroup.com'),
    ('Muhammad',         'Faheem',          '2107', 'muhammad.faheem@4bgroup.com'),
    ('Asad',             'Ullah',           '2115', 'asadullah.khan@4bgroup.com'),
    ('Ammar',            'Naqvi',           '2122', 'ammar.naqvi@4bgroup.com'),
    ('Nazir',            'Ahmad',           '2127', 'nazir.ahmad@4bgroup.com'),
    ('Muhammad',         'Dawood',          '2129', 'muhammd.dawood@4bgroup.com'),
    ('Sajan',            'Das',             '2132', 'sajan.das@4bgroup.com'),
    ('Maaz',             'Haider',          '2137', 'maaz.haider@4bgroup.com'),
    ('Tariq',            'Mehmood',         '2151', 'tariq.jatoi@4bgroup.com'),
    ('Nadir',            'Hussain',         '2152', 'nadir.hussain@4bgroup.com'),
    ('Pervaiz',          'Iqbal',           '2155', 'pervaiz.iqbal@4bgroup.com'),
    ('Shafqat',          'Abbas',           '2164', 'shafqat.abbas@4bgroup.com'),
    ('Amjad',            'Ali',             '2166', 'amjad.ali@4bgroup.com'),
    ('Muhammad',         'Jamshed',         '2180', 'muhammad.jamshed@4bgroup.com'),
    ('Muhammad',         'Sadaqat',         '2186', 'muhammad.sadaqat@4bgroup.com'),
    ('Umer',             'Hayat',           '2224', 'umer.hayat@4bgroup.com'),
    ('Muhammad',         'Usman',           '2227', 'muhammad.usman@4bgroup.com'),
    ('Pervez',           'Ali',             '2233', 'pervez.ali@4bgroup.com'),
    ('Ghulam',           'Muhammad',        '2234', 'ghulam.muhammad@4bgroup.com'),
    ('Ghulam',           'Muhammad',        '2235', 'ghulam.muhammad@gmail.com'),
    ('Kifayat',          'Hussain',         '2238', 'kifayat.jatoi@4bgroup.com'),
    ('Javed',            'Dhiloo',          '2239', 'javed.dhiloo@4bgroup.com'),
    ('Azhar',            'Hussain',         '2244', 'azhar.hussain@4bgroup.com'),
    ('Nusrat',           'Hussain',         '2252', 'nusrat.hussain@4bgroup.com'),
    ('Muhammad',         'Fayaz',           '2457', 'fayaz.rasheed@4bgroup.com'),
]

role = Role.objects.get(id=SALES_STAFF_ROLE_ID)
results = []
skipped = []

with transaction.atomic():
    for first_name, last_name, emp_code, email in USERS:
        # Skip if email already exists
        if User.objects.filter(email__iexact=email).exists():
            skipped.append((emp_code, email, 'email already exists'))
            continue
        # Skip if employee code already assigned
        if SalesStaffProfile.objects.filter(employee_code=emp_code).exists():
            skipped.append((emp_code, email, 'employee code already exists'))
            continue

        username = email.split('@')[0].lower()
        # Ensure username uniqueness
        base_username = username
        counter = 1
        while User.objects.filter(username=username).exists():
            username = f"{base_username}{counter}"
            counter += 1

        plain_password = build_password(first_name)

        user = User(
            email=email,
            username=username,
            first_name=first_name,
            last_name=last_name,
            role=role,
            is_active=True,
            is_sales_staff=True,
            is_staff=False,
        )
        user.set_password(plain_password)
        user.save()

        profile = SalesStaffProfile.objects.create(
            user=user,
            employee_code=emp_code,
        )

        results.append({
            'Employee Code': emp_code,
            'First Name': first_name,
            'Last Name': last_name,
            'Email': email,
            'Username': username,
            'Password': plain_password,
            'Role': 'Sales Staff',
        })
        print(f"  CREATED  {emp_code}  {email}  pwd={plain_password}")

print(f"\nCreated: {len(results)}  |  Skipped: {len(skipped)}")
for s in skipped:
    print(f"  SKIPPED  {s[0]}  {s[1]}  ({s[2]})")

# ── Export to XLSX ─────────────────────────────────────────────────────────────
if results:
    wb = Workbook()
    ws = wb.active
    ws.title = "New Sales Staff"

    headers = ['Employee Code', 'First Name', 'Last Name', 'Email', 'Username', 'Password', 'Role']
    hdr_font  = Font(bold=True, color="FFFFFF")
    hdr_fill  = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    hdr_align = Alignment(horizontal="center", vertical="center")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = hdr_align

    for row, r in enumerate(results, 2):
        for col, h in enumerate(headers, 1):
            ws.cell(row=row, column=col, value=r[h])

    # Auto column width
    for col in ws.columns:
        max_len = max((len(str(c.value or '')) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    out_path = '/home/www/django_web_portal/new_sales_staff_passwords.xlsx'
    wb.save(out_path)
    print(f"\nXLSX saved → {out_path}")
else:
    print("Nothing to export.")
