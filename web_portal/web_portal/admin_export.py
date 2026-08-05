"""
Site-wide admin export actions (CSV + Excel).

Registered once on the custom admin site (see web_portal/admin.py), so every
model changelist gains "Export selected -> CSV / Excel" without touching each
of the ~60 ModelAdmin classes.

Column selection, in priority order:
  1. `export_fields` attribute on the ModelAdmin (list of field/callable names)
  2. the visible `list_display` (minus the selection checkbox)
  3. the model's own concrete fields
Values are resolved exactly the way the changelist resolves them, then any
HTML from badge/image callables is stripped so the cell holds plain text.
"""
import csv
import datetime
import html
import re
from io import BytesIO

from django.contrib import admin
from django.contrib.admin.utils import label_for_field, lookup_field
from django.http import HttpResponse
from django.utils import timezone
from django.utils.encoding import force_str
from django.utils.html import strip_tags

_SHEET_TITLE_BAD = re.compile(r'[\[\]:*?/\\]')


def _export_columns(model_admin):
    fields = getattr(model_admin, 'export_fields', None)
    if fields:
        return list(fields)
    cols = [c for c in (model_admin.list_display or ()) if c != 'action_checkbox']
    if not cols or cols == ['__str__']:
        return [f.name for f in model_admin.model._meta.fields]
    return cols


def _header(model_admin, name):
    try:
        label = force_str(label_for_field(name, model_admin.model, model_admin)).strip()
    except Exception:
        label = force_str(name).replace('_', ' ').strip().title()
    # Field labels arrive lower case ("check in time"), while a callable's
    # description is written properly ("Employee name"), so a sheet mixed both.
    # Lift only the first character - .title() would mangle acronyms like RSM.
    return label[:1].upper() + label[1:] if label else label


def _cell(model_admin, obj, name):
    try:
        _field, _attr, value = lookup_field(name, obj, model_admin)
    except Exception:
        value = getattr(obj, name, '')
        if callable(value):
            try:
                value = value()
            except Exception:
                value = ''
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'Yes' if value else 'No'
    # Datetimes reach here raw, so without this a cell reads
    # "2026-08-05 02:13:47.423000+00:00" - UTC, to the microsecond.
    if isinstance(value, datetime.datetime):
        if timezone.is_aware(value):
            value = timezone.localtime(value)
        return value.strftime('%Y-%m-%d %H:%M')
    if isinstance(value, datetime.date):
        return value.strftime('%Y-%m-%d')
    text = force_str(value)
    if '<' in text and '>' in text:          # HTML from a badge / image callable
        text = strip_tags(text)
    # Placeholders like the "&mdash;" an empty thumbnail returns carry no tags,
    # so strip_tags leaves them as literal entity text in the sheet.
    return html.unescape(text).strip()


def _rows(model_admin, queryset, columns):
    yield [_header(model_admin, c) for c in columns]
    for obj in queryset:
        yield [_cell(model_admin, obj, c) for c in columns]


def _filename(model_admin, ext):
    stamp = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    return f'{model_admin.model._meta.model_name}_export_{stamp}.{ext}'


def _get_summary(model_admin, request):
    """
    Optional summary block prepended to the export.

    A ModelAdmin may define `get_export_summary(request)` returning:
        {'title': 'Attendance summary — 2026-07-24',
         'rows': [('Present today', 0), ('Off today', 11), ...]}
    Returns None when the admin provides nothing.
    """
    fn = getattr(model_admin, 'get_export_summary', None)
    if not fn:
        return None
    try:
        data = fn(request)
    except Exception:
        return None
    if not data or not (data.get('title') or data.get('rows')):
        return None
    return data


@admin.action(description='Export selected → CSV')
def export_as_csv(model_admin, request, queryset):
    columns = _export_columns(model_admin)
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = (
        f'attachment; filename="{_filename(model_admin, "csv")}"'
    )
    response.write('﻿')                 # UTF-8 BOM so Excel reads it correctly
    writer = csv.writer(response)

    summary = _get_summary(model_admin, request)
    if summary:
        if summary.get('title'):
            writer.writerow([summary['title']])
        for label, value in summary.get('rows', []):
            writer.writerow([label, value])
        writer.writerow([])            # blank spacer before the data table

    for row in _rows(model_admin, queryset, columns):
        writer.writerow(row)
    return response


@admin.action(description='Export selected → Excel (.xlsx)')
def export_as_xlsx(model_admin, request, queryset):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    columns = _export_columns(model_admin)
    wb = Workbook()
    ws = wb.active
    title = _SHEET_TITLE_BAD.sub('', str(model_admin.model._meta.verbose_name_plural))
    ws.title = (title or 'Export')[:31]

    header_fill = PatternFill('solid', fgColor='E26830')
    header_font = Font(bold=True, color='FFFFFF')

    # Optional summary block at the top.
    r_idx = 1
    summary = _get_summary(model_admin, request)
    if summary:
        if summary.get('title'):
            cell = ws.cell(r_idx, 1, summary['title'])
            cell.font = Font(bold=True, size=13, color='1F4E78')
            r_idx += 1
        for label, value in summary.get('rows', []):
            ws.cell(r_idx, 1, label).font = Font(bold=True)
            ws.cell(r_idx, 2, value)
            r_idx += 1
        r_idx += 1                     # blank spacer row

    header_row = r_idx
    widths = {}
    for row in _rows(model_admin, queryset, columns):
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(r_idx, c_idx, value)
            if r_idx == header_row:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            # Floor of 16 so short columns are still comfortably readable rather
            # than clipped to their header; ceiling so one long cell (a manager's
            # whole territory list) cannot push the others off the screen.
            widths[c_idx] = min(max(widths.get(c_idx, 16), len(str(value)) + 3), 60)
        r_idx += 1

    for c_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(c_idx)].width = width
    ws.freeze_panes = ws.cell(header_row + 1, 1)
    if columns and r_idx - 1 >= header_row:
        ws.auto_filter.ref = f'A{header_row}:{get_column_letter(len(columns))}{r_idx - 1}'

    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(
        buf.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = (
        f'attachment; filename="{_filename(model_admin, "xlsx")}"'
    )
    return response
